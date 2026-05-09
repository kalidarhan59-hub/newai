"""Universal file parser — PDF / DOCX / XLSX / CSV / JSON / TXT / ZIP / code."""

from __future__ import annotations

import csv
import io
import json
import logging
import zipfile
from pathlib import Path
from typing import Any

from ..schemas import ParsedDocument

logger = logging.getLogger(__name__)

CODE_EXTS = {
    ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".cc", ".cpp", ".h",
    ".hpp", ".rs", ".go", ".rb", ".php", ".sh", ".sql", ".swift", ".kt",
    ".html", ".css", ".scss", ".yaml", ".yml", ".toml",
}
TEXT_EXTS = {".txt", ".md", ".rst", ".log"}


def parse_path(path: Path, *, name: str | None = None) -> ParsedDocument:
    return parse_bytes(path.read_bytes(), name or path.name)


def parse_bytes(data: bytes, name: str) -> ParsedDocument:
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            return _parse_pdf(data, name)
        if suffix == ".docx":
            return _parse_docx(data, name)
        if suffix == ".xlsx":
            return _parse_xlsx(data, name)
        if suffix == ".csv":
            return _parse_csv(data, name)
        if suffix == ".json":
            return _parse_json(data, name)
        if suffix == ".zip":
            return _parse_zip(data, name)
        if suffix in TEXT_EXTS or suffix in CODE_EXTS:
            return _parse_text(data, name, suffix)
        # Best-effort UTF-8 decode for unknown types.
        return _parse_text(data, name, suffix or ".bin")
    except Exception as e:
        logger.exception("File parse failed for %s: %s", name, e)
        return ParsedDocument(
            name=name,
            content_type=suffix or "application/octet-stream",
            summary=f"Could not parse {name}: {e}",
            chunks=[],
            metadata={"error": str(e)},
        )


# ---------------------------------------------------------------------------
# Per-format parsers
# ---------------------------------------------------------------------------


def _parse_pdf(data: bytes, name: str) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n\n".join(pages)
    return _build(name, "application/pdf", text, {"pages": len(pages)})


def _parse_docx(data: bytes, name: str) -> ParsedDocument:
    from docx import Document

    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs if p.text)
    return _build(
        name,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        text,
        {"paragraphs": len(doc.paragraphs)},
    )


def _parse_xlsx(data: bytes, name: str) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: list[str] = []
    rows = 0
    for ws in wb.worksheets:
        sheet_lines = [f"# {ws.title}"]
        for row in ws.iter_rows(values_only=True):
            sheet_lines.append("\t".join("" if c is None else str(c) for c in row))
            rows += 1
        sheets.append("\n".join(sheet_lines))
    text = "\n\n".join(sheets)
    return _build(
        name,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        text,
        {"sheets": len(wb.worksheets), "rows": rows},
    )


def _parse_csv(data: bytes, name: str) -> ParsedDocument:
    text_io = io.StringIO(data.decode("utf-8", errors="replace"))
    reader = csv.reader(text_io)
    rows = list(reader)
    text = "\n".join("\t".join(r) for r in rows)
    return _build(name, "text/csv", text, {"rows": len(rows)})


def _parse_json(data: bytes, name: str) -> ParsedDocument:
    parsed = json.loads(data.decode("utf-8", errors="replace"))
    text = json.dumps(parsed, indent=2, ensure_ascii=False)
    return _build(name, "application/json", text, {"top_level_keys": _json_keys(parsed)})


def _parse_zip(data: bytes, name: str) -> ParsedDocument:
    summaries: list[str] = []
    chunks: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            if info.is_dir() or info.file_size > 1_000_000:
                continue
            try:
                inner = zf.read(info.filename)
            except RuntimeError:
                continue
            doc = parse_bytes(inner, info.filename)
            summaries.append(f"- {info.filename}: {doc.summary}")
            chunks.extend(f"# {info.filename}\n{c}" for c in doc.chunks[:3])
    summary = "ZIP archive contents:\n" + "\n".join(summaries[:50])
    return ParsedDocument(
        name=name,
        content_type="application/zip",
        summary=summary,
        chunks=chunks[:50],
        metadata={"entries": len(summaries)},
    )


def _parse_text(data: bytes, name: str, suffix: str) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    return _build(name, f"text/{suffix.lstrip('.') or 'plain'}", text, {})


# ---------------------------------------------------------------------------
# Utilities
# ---------------------------------------------------------------------------


def _build(name: str, content_type: str, text: str, metadata: dict[str, Any]) -> ParsedDocument:
    chunks = _chunk(text)
    summary = _summary(text)
    return ParsedDocument(
        name=name,
        content_type=content_type,
        summary=summary,
        chunks=chunks,
        metadata=metadata,
    )


def _chunk(text: str, *, max_chars: int = 800) -> list[str]:
    text = text.strip()
    if not text:
        return []
    out: list[str] = []
    i = 0
    while i < len(text):
        out.append(text[i : i + max_chars])
        i += max_chars
    return out[:200]


def _summary(text: str, *, limit: int = 320) -> str:
    text = " ".join(text.split())
    if not text:
        return "(empty file)"
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _json_keys(parsed: Any) -> list[str]:
    if isinstance(parsed, dict):
        return list(parsed.keys())[:25]
    if isinstance(parsed, list) and parsed and isinstance(parsed[0], dict):
        return list(parsed[0].keys())[:25]
    return []
